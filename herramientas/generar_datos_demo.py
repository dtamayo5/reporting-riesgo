"""
Genera todos los datos de la demostración (entidad y emisores ficticios):

  mercado/factores_mercado.csv          historia sintética de factores de riesgo
  entrada/politica_riesgo.xlsx          límites, emisores, haircuts, escenarios
  entrada/posiciones_AAAA-MM-DD.xlsx    seis cortes mensuales (abr-sep 2026)
  entrada/plantilla/plantilla_posiciones.xlsx   formato vacío para un cliente real

Los Excel de posiciones imitan a propósito un archivo manual real: títulos
arriba, fila de totales, un emisor escrito de varias formas, fechas como
texto, una tasa en porcentaje, un duplicado, un título vencido y un
instrumento no soportado. El motor debe detectarlo todo.

Uso:  python herramientas/generar_datos_demo.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "herramientas"))

import mercado_sintetico  # noqa: E402
from riesgo import config as C  # noqa: E402
from riesgo.lectura import (calificacion_lp, leer_politica, normalizar_calificacion,  # noqa: E402
                            normalizar_frecuencia, normalizar_tipo)
from riesgo.liquidez import calcular_liquidez  # noqa: E402
from riesgo.mercado import Mercado  # noqa: E402
from riesgo.valoracion import Portafolio  # noqa: E402

ENTIDAD = "Cooperativa Financiera Ejemplo"
CORTES = ["2026-04-30", "2026-05-29", "2026-06-30", "2026-07-31", "2026-08-31", "2026-09-30"]
PATRIMONIO = [96_800, 97_100, 97_500, 97_900, 98_200, 98_450]           # COP millones
IRL30_OBJETIVO = [1.86, 1.77, 1.65, 1.54, 1.43, 1.29]
SEP, AGO = 5, 4

FUENTE = "Arial"
F_TITULO = Font(name=FUENTE, bold=True, size=12)
F_SUB = Font(name=FUENTE, italic=True, size=9, color="52514E")
F_ENC = Font(name=FUENTE, bold=True, size=9, color="FFFFFF")
F_NORMAL = Font(name=FUENTE, size=9)
F_NEGRITA = Font(name=FUENTE, size=9, bold=True)
RELLENO_ENC = PatternFill("solid", fgColor="1F3B5C")
RELLENO_EJEMPLO = PatternFill("solid", fgColor="FFF2CC")
BORDE = Border(bottom=Side(style="thin", color="C9D1DC"))
FMT_MONTO, FMT_PCT, FMT_FECHA = "#,##0", "0.00%", "dd/mm/yyyy"

# ----------------------------------------------------------------------------
# Portafolio de la entidad ficticia. Montos en COP millones (USD en dólares,
# ETF en unidades). 'desde'/'hasta' = índice del primer y último corte.
# 'errores' = defectos inyectados a propósito en el Excel de ese corte.
# ----------------------------------------------------------------------------
POSICIONES = [
    dict(id=101, especie="Cuenta de ahorros 2201", tipo="Cuenta de ahorros", emisor="Banco Altiplano",
         cal="AAA", port="Tesorería", clas="Disponible", saldo=[3500, 3300, 3600, 3400, 3200, 3000]),
    dict(id=102, especie="Cuenta de ahorros 4410", tipo="Cuenta de ahorros", emisor="Banco Cordillera",
         cal="AAA", port="Tesorería", clas="Disponible", saldo=[5000, 5400, 6000, 6300, 6500, 7200]),
    dict(id=103, especie="FIC Renta Vista", tipo="FIC vista", emisor="FIC Renta Vista (Fiduciaria Sabana)",
         cal="AAA", port="Tesorería", clas="Disponible", saldo=[5200, 5600, 5100, 5400, 5800, 5500]),
    dict(id=201, especie="TES TF 5,75% nov-2027", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Fondo de liquidez", clas="Disponible para la venta", nom=9000, tasa=0.0575, per="AV",
         emi="2020-11-03", venc="2027-11-03"),
    dict(id=202, especie="TES TF 6,00% abr-2028", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Negociable", nom=8000, tasa=0.06, per="AV", emi="2012-04-28", venc="2028-04-28"),
    dict(id=203, especie="TES TF 7,75% sep-2030", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Negociable", nom=7000, tasa=0.0775, per="AV", emi="2014-09-18", venc="2030-09-18"),
    dict(id=204, especie="TES TF 7,00% mar-2031", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Disponible para la venta", nom=6000, tasa=0.07, per="AV", emi="2021-03-26",
         venc="2031-03-26"),
    dict(id=205, especie="TES TF 7,00% jun-2032", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Disponible para la venta", nom=5000, tasa=0.07, per="AV", emi="2016-06-30",
         venc="2032-06-30", restr="Sí"),
    dict(id=206, especie="TES TF 13,25% feb-2033", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Negociable", nom=4000, tasa=0.1325, per="AV", emi="2023-02-09", venc="2033-02-09"),
    dict(id=207, especie="TES TF 7,25% oct-2034", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Negociable", nom=6000, tasa=0.0725, per="AV", emi="2019-10-18",
         venc="2034-10-18", desde=3),
    dict(id=208, especie="TES TF 6,25% jul-2036", tipo="TES", emisor="Nación - MHCP", cal="Nación",
         port="Tesorería", clas="Negociable", nom=3000, tasa=0.0625, per="AV", emi="2021-07-09", venc="2036-07-09"),
    dict(id=301, especie="CDT Banco Cordillera 10,90% TV", tipo="CDT", emisor="Banco Cordillera S.A.",
         cal="AAA", port="Tesorería", clas="Al vencimiento", nom=4500, tasa=0.109, per="TV",
         emi="2025-06-15", venc="2027-06-15"),
    dict(id=302, especie="CDT Banco Cordillera 11,40% SV", tipo="CDT", emisor="BANCO CORDILLERA",
         cal="AAA", port="Tesorería", clas="Disponible para la venta", nom=3500, tasa=0.114, per="SV",
         emi="2026-03-20", venc="2028-03-20", desde=2),
    dict(id=303, especie="CDT Banco Cordillera 12,10% AV", tipo="CDT", emisor="Bco. Cordillera",
         cal="AAA", port="Tesorería", clas="Disponible para la venta", nom=3000, tasa=0.121, per="AV",
         emi="2025-12-15", venc="2027-12-15", desde=4,
         errores={AGO: ["venc_texto", "tasa_pct"], SEP: ["venc_texto"]}),
    dict(id=304, especie="CDT IBR Banco Altiplano IBR+1,90% TV", tipo="CDT IBR", emisor="Banco Altiplano",
         cal="AAA", port="Tesorería", clas="Disponible para la venta", nom=4000, margen=0.019, per="TV",
         emi="2025-09-10", venc="2028-09-10"),
    dict(id=305, especie="CDT IBR Banco Guayacán IBR+2,60% TV", tipo="CDT IBR", emisor="Banco Guayacán",
         cal="AA+", port="Tesorería", clas="Disponible para la venta", nom=4000, margen=0.026, per="TV",
         emi="2025-11-05", venc="2027-11-05", errores={SEP: ["sin_calificacion"]}),
    dict(id=306, especie="CDT Financiera Río Claro 12,80% SV", tipo="CDT", emisor="Financiera Río Claro",
         cal="AA", port="Tesorería", clas="Al vencimiento", nom=3000, tasa=0.128, per="SV",
         emi="2025-10-22", venc="2027-04-22", errores={SEP: ["tasa_pct"]}),
    dict(id=307, especie="CDT Banco Guayacán 11,90% PV", tipo="CDT", emisor="Banco Guayacán",
         cal="AA+", port="Tesorería", clas="Al vencimiento", nom=2500, tasa=0.119, per="PV",
         emi="2026-03-10", venc="2027-03-10", errores={SEP: ["calificacion_distinta"]}),
    dict(id=308, especie="CDT Banco Altiplano 11,60% PV", tipo="CDT", emisor="Banco Altiplano",
         cal="F1+", port="Tesorería", clas="Al vencimiento", nom=2500, tasa=0.116, per="PV",
         emi="2026-06-01", venc="2026-12-01", desde=2),
    dict(id=309, especie="CDT Leasing Nevado 12,40% TV", tipo="CDT", emisor="Leasing Nevado",
         cal="AA", port="Tesorería", clas="Disponible para la venta", nom=2000, tasa=0.124, per="TV",
         emi="2025-09-30", venc="2027-09-30", errores={SEP: ["nominal_texto"]}),
    dict(id=310, especie="CDT Leasing Nevado 10,20% PV", tipo="CDT", emisor="Leasing Nevado S.A.",
         cal="AA", port="Tesorería", clas="Al vencimiento", nom=1500, tasa=0.102, per="PV",
         emi="2025-09-15", venc="2026-09-15", errores={SEP: ["vencido"]}),
    dict(id=401, especie="Bono Energía del Norte 11,50% TV 2031", tipo="Bono ordinario",
         emisor="Energía del Norte S.A. E.S.P.", cal="AAA", port="Tesorería", clas="Disponible para la venta",
         nom=4000, tasa=0.115, per="TV", emi="2024-05-15", venc="2031-05-15", errores={SEP: ["duplicado"]}),
    dict(id=402, especie="Bono Corporación Sabana 12,20% SV 2029", tipo="Bono ordinario",
         emisor="Corporación Sabana", cal="AA+", port="Tesorería", clas="Disponible para la venta",
         nom=3500, tasa=0.122, per="SV", emi="2024-08-30", venc="2029-08-30", errores={SEP: ["valor_distinto"]}),
    dict(id=403, especie="Bono Constructora Ceiba 14,80% TV 2028", tipo="Bono ordinario",
         emisor="Constructora Ceiba S.A.", cal="A+", port="Tesorería", clas="Disponible para la venta",
         nom=2500, tasa=0.148, per="TV", emi="2025-11-30", venc="2028-11-30"),
    dict(id=501, especie="ETF acciones Colombia", tipo="ETF acciones locales", emisor="ETF Índice Local",
         cal="", port="Tesorería", clas="Negociable", unidades=1_500_000, beta=1.0),
    dict(id=601, especie="Cuenta de compensación USD", tipo="Depósito en USD", emisor="Banco Guayacán",
         cal="AA+", port="Tesorería", clas="Disponible", usd=950_000),
    dict(id=701, especie="Titularización hipotecaria TIPS E-17", tipo="Titularización TIPS",
         emisor="Titularizadora Ejemplo", cal="AAA", port="Tesorería", clas="Al vencimiento", nom=1200,
         tasa=0.105, per="MV", emi="2023-05-01", venc="2030-05-01", desde=5),
]

MAESTRO = [
    ("Nación - MHCP", "Nación; MHCP; Ministerio de Hacienda; Gobierno Nacional", "Gobierno", "Nación", "Sí"),
    ("Banco Cordillera", "", "Financiero", "AAA", "No"),
    ("Banco Altiplano", "Altiplano", "Financiero", "AAA", "No"),
    ("Banco Guayacán", "Guayacan", "Financiero", "AA+", "No"),
    ("Financiera Río Claro", "Rio Claro", "Financiero", "AA", "No"),
    ("Leasing Nevado", "", "Financiero", "AA", "No"),
    ("Energía del Norte S.A. E.S.P.", "Energía del Norte", "Real", "AAA", "No"),
    ("Corporación Sabana", "Sabana", "Real", "AA+", "No"),
    ("Constructora Ceiba", "", "Real", "A+", "No"),
    ("FIC Renta Vista (Fiduciaria Sabana)", "FIC Renta Vista", "Fondos", "AAA", "No"),
    ("ETF Índice Local", "", "Renta variable", "", "No"),
    ("Titularizadora Ejemplo", "", "Financiero", "AAA", "No"),
]

LIMITES = [  # código, indicador, categoría, tipo, límite, alerta, unidad
    ("VAR_HS_10D_PCT_PT", "VaR histórico 99% a 10 días / patrimonio técnico", "Mercado", "Máximo", 0.015, 0.012, "%"),
    ("DURACION_RF", "Duración modificada de la renta fija", "Mercado", "Máximo", 2.5, 2.2, "años"),
    ("DV01_TOTAL", "Pérdida por +1 pb en toda la curva TES", "Mercado", "Máximo", 25.0, 20.0, "COP M"),
    ("ESTRES_PEOR_PCT_PT", "Peor escenario de estrés / patrimonio técnico", "Mercado", "Máximo", 0.08, 0.065, "%"),
    ("BT_EXCEPCIONES_HS", "Excepciones del backtesting (VaR histórico, 250 días)", "Mercado", "Máximo", 9, 4, "número"),
    ("USD_PCT_PT", "Posición en moneda extranjera / patrimonio técnico", "Mercado", "Máximo", 0.05, 0.04, "%"),
    ("RV_PCT", "Renta variable / portafolio", "Mercado", "Máximo", 0.05, 0.04, "%"),
    ("EMISOR_MAX_PCT", "Mayor exposición a un emisor privado / portafolio", "Concentración", "Máximo", 0.15, 0.12, "%"),
    ("HHI_PRIVADOS", "Índice Herfindahl-Hirschman de emisores privados", "Concentración", "Máximo", 2500, 1800, "puntos"),
    ("TOP5_PCT", "Cinco mayores emisores privados / portafolio", "Concentración", "Máximo", 0.55, 0.50, "%"),
    ("CALIF_A_O_INF_PCT", "Renta fija con calificación A+ o inferior / portafolio", "Concentración", "Máximo", 0.05, 0.04, "%"),
    ("IRL_7D", "IRL razón a 7 días", "Liquidez", "Mínimo", 1.20, 1.50, "%"),
    ("IRL_30D", "IRL razón a 30 días", "Liquidez", "Mínimo", 1.10, 1.40, "%"),
    ("ERRORES_DATOS", "Filas excluidas por errores de datos", "Calidad de datos", "Máximo", 3, 0, "número"),
]

HAIRCUTS = [
    ("LIQUIDEZ", "Disponible y FIC a la vista", 0.0),
    ("TES_HASTA_1A", "TES con vencimiento hasta 1 año", 0.02),
    ("TES_1A_5A", "TES entre 1 y 5 años", 0.05),
    ("TES_MAS_5A", "TES a más de 5 años", 0.10),
    ("PRIVADO_AAA", "CDT y bonos AAA (o F1+ / BRC 1+)", 0.15),
    ("PRIVADO_AA", "CDT y bonos AA+ a AA-", 0.25),
    ("PRIVADO_INFERIOR", "CDT y bonos A+ o inferior", 0.50),
    ("RENTA_VARIABLE", "Renta variable", 0.40),
    ("DIVISA", "Posiciones en moneda extranjera", 0.10),
]

PARAM_LIQUIDEZ = [
    ("1-7 días", 0.030, "Retiro estimado de depósitos a la vista en la banda"),
    ("8-15 días", 0.020, "Retiro estimado de depósitos a la vista en la banda"),
    ("16-30 días", 0.035, "Retiro estimado de depósitos a la vista en la banda"),
    ("31-90 días", 0.060, "Retiro estimado de depósitos a la vista en la banda"),
    ("Tope de entradas", 0.75, "Las entradas computan hasta este % de las salidas de la banda"),
]

ESCENARIOS = [  # código, nombre, tipo, paralelo, corto, largo, AAA, AA, TRM, RV, desde, hasta, descripción
    ("E1", "Alza paralela de 200 pb", "Hipotético", 200, 0, 0, 0, 0, 0, 0, None, None,
     "Toda la curva TES sube 200 pb."),
    ("E2", "Empinamiento de la curva", "Hipotético", 0, -50, 150, 0, 0, 0, 0, None, None,
     "Tramo corto -50 pb y tramo largo +150 pb."),
    ("E3", "Ampliación de spreads", "Hipotético", 0, 0, 0, 100, 200, 0, 0, None, None,
     "Spreads AAA +100 pb y AA +200 pb."),
    ("E4", "Revaluación del peso", "Hipotético", 0, 0, 0, 0, 0, -0.15, 0, None, None,
     "La TRM cae 15% (pierde la posición en dólares)."),
    ("E5", "Estrés local combinado", "Hipotético", 200, 0, 0, 100, 200, 0.15, -0.25, None, None,
     "Tasas +200 pb, spreads AAA +100/AA +200, TRM +15%, renta variable -25%."),
    ("E6", "Caída paralela de 100 pb", "Hipotético", -100, 0, 0, 0, 0, 0, 0, None, None,
     "Escenario favorable, informativo."),
    ("H1", "Réplica: venta masiva de TES 2022", "Histórico", 0, 0, 0, 0, 0, 0, 0, "2022-09-01", "2022-12-31",
     "Peor ventana de 20 días hábiles del periodo aplicada al portafolio actual."),
    ("H2", "Réplica: alzas de tasas de 2026", "Histórico", 0, 0, 0, 0, 0, 0, 0, "2026-01-01", "2026-04-30",
     "Peor ventana de 20 días hábiles del periodo aplicada al portafolio actual."),
]


# ------------------------------------------------------------ utilidades Excel
def _titulos(ws, titulo, subtitulo, extra=None):
    ws["A1"] = titulo
    ws["A1"].font = F_TITULO
    ws["A2"] = subtitulo
    ws["A2"].font = F_SUB
    if extra:
        ws["A3"] = extra
        ws["A3"].font = F_SUB


def _encabezados(ws, fila, nombres, anchos):
    for j, (nombre, ancho) in enumerate(zip(nombres, anchos), start=1):
        c = ws.cell(fila, j, nombre)
        c.font, c.fill = F_ENC, RELLENO_ENC
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila].height = 30


def _fila(ws, fila, valores, formatos, negrita=False):
    for j, (v, fmt) in enumerate(zip(valores, formatos), start=1):
        c = ws.cell(fila, j, v)
        c.font = F_NEGRITA if negrita else F_NORMAL
        c.border = BORDE
        if fmt:
            c.number_format = fmt


def _libro():
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    return wb


# --------------------------------------------------------------- la política
def escribir_politica(ruta: Path):
    wb = _libro()
    ws = wb.active
    ws.title = "Entidad"
    _titulos(ws, "Política de riesgos de tesorería", f"{ENTIDAD} (entidad ficticia de demostración)",
             "Parámetros de negocio del motor. Los aprueba el Consejo; el motor solo los aplica.")
    _encabezados(ws, 5, ["Parámetro", "Valor", "Nota"], [38, 34, 60])
    filas = [("Nombre de la entidad", ENTIDAD, "Aparece en la portada del informe."),
             ("Tipo de entidad", "Cooperativa financiera", ""),
             ("Órgano que aprueba", "Consejo de Administración", ""),
             ("Fecha de aprobación de la política", "26/02/2026", "Ilustrativa."),
             ("Moneda de reporte", "COP millones", ""),
             ("Modo demostración", "Sí", "Con 'Sí' el informe aclara que los datos son de demostración."),
             ("Nota de pie de página", "Entidad y emisores ficticios; datos de mercado sintéticos.",
              "Texto opcional al pie de cada página del informe.")]
    for i, f in enumerate(filas):
        _fila(ws, 6 + i, f, [None, None, None])

    ws = wb.create_sheet("Limites")
    _titulos(ws, "Límites y umbrales de alerta",
             "Límites ILUSTRATIVOS: reemplazar por los aprobados en la política de la entidad.",
             "Máximo: rojo si valor > límite; ámbar si valor > alerta. Mínimo: rojo si valor < límite; "
             "ámbar si valor < alerta. Porcentajes guardados como fracción (15% = 0,15).")
    _encabezados(ws, 5, ["Código", "Indicador", "Categoría", "Tipo de límite", "Límite", "Alerta", "Unidad"],
                 [22, 52, 16, 14, 11, 11, 10])
    for i, (cod, ind, cat, tipo, lim, ale, uni) in enumerate(LIMITES):
        fmt = FMT_PCT if uni == "%" else "#,##0.0"
        _fila(ws, 6 + i, [cod, ind, cat, tipo, lim, ale, uni], [None, None, None, None, fmt, fmt, None])

    ws = wb.create_sheet("Emisores")
    _titulos(ws, "Maestro de emisores",
             "Un emisor por fila. En 'Alias' separe los otros nombres con punto y coma (;).",
             "Las diferencias de mayúsculas, tildes, puntuación, 'S.A.' y abreviaturas como 'Bco.' "
             "se resuelven solas; los alias cubren nombres realmente distintos.")
    _encabezados(ws, 5, ["Emisor oficial", "Alias", "Sector", "Calificación", "Exento de límite"],
                 [38, 52, 16, 13, 16])
    for i, f in enumerate(MAESTRO):
        _fila(ws, 6 + i, f, [None] * 5)

    ws = wb.create_sheet("Haircuts")
    _titulos(ws, "Descuentos (haircuts) de activos líquidos",
             "Ilustrativos, inspirados en la lógica del IRL regulatorio. Fracción: 5% = 0,05.")
    _encabezados(ws, 4, ["Categoría", "Descripción", "Haircut"], [20, 44, 11])
    for i, f in enumerate(HAIRCUTS):
        _fila(ws, 5 + i, f, [None, None, FMT_PCT])

    ws = wb.create_sheet("Liquidez")
    _titulos(ws, "Parámetros de liquidez", "Factores de retiro de depósitos a la vista por banda "
             "(ilustrativos: se calibran con la historia de retiros de la entidad).")
    _encabezados(ws, 4, ["Parámetro", "Valor", "Nota"], [20, 11, 62])
    for i, f in enumerate(PARAM_LIQUIDEZ):
        _fila(ws, 5 + i, f, [None, FMT_PCT, None])

    ws = wb.create_sheet("Escenarios_estres")
    _titulos(ws, "Escenarios de estrés", "Hipotéticos: choques instantáneos. Históricos: peor ventana de "
             "20 días hábiles dentro del periodo 'Desde'-'Hasta', aplicada al portafolio actual.")
    _encabezados(ws, 4, ["Código", "Nombre", "Tipo", "Paralelo pb", "Corto pb", "Largo pb",
                         "Spread AAA pb", "Spread AA pb", "TRM %", "Renta variable %", "Desde",
                         "Hasta", "Descripción"],
                 [8, 34, 11, 10, 9, 9, 11, 11, 9, 12, 12, 12, 60])
    for i, e in enumerate(ESCENARIOS):
        vals = list(e)
        vals[10] = pd.Timestamp(vals[10]).to_pydatetime() if vals[10] else None
        vals[11] = pd.Timestamp(vals[11]).to_pydatetime() if vals[11] else None
        _fila(ws, 5 + i, vals, [None, None, None, "0", "0", "0", "0", "0", "0%", "0%",
                                FMT_FECHA, FMT_FECHA, None])
    wb.save(ruta)


# ------------------------------------------------------- posiciones del corte
def _canonica(p: dict, corte: pd.Timestamp) -> dict:
    """Fila limpia (formato interno del motor) para valorar en el generador."""
    tipo, _ = normalizar_tipo(p["tipo"])
    cal = "NACION" if tipo == "TES_TF" else normalizar_calificacion(p.get("cal")) or None
    return {
        "tipo": tipo, "valor_nominal": p.get("nom", 0) * 1e6 if "nom" in p else p.get("usd", np.nan),
        "tasa_facial": p.get("tasa", np.nan), "margen": p.get("margen", np.nan), "base_tasa": "EA",
        "frecuencia": normalizar_frecuencia(p.get("per")) if p.get("per") else 0,
        "fecha_emision": pd.Timestamp(p["emi"]) if "emi" in p else pd.NaT,
        "fecha_vencimiento": pd.Timestamp(p["venc"]) if "venc" in p else pd.NaT,
        "valor_mercado": np.nan, "beta": p.get("beta", 1.0),
        "calificacion_lp": calificacion_lp(cal) if cal else None,
        "restringido": p.get("restr") == "Sí",
    }


def _vigente(p, k, corte):
    if k < p.get("desde", 0) or k > p.get("hasta", 99):
        return False
    return True


def construir_corte(k: int, mercado: Mercado, rng):
    corte = pd.Timestamp(CORTES[k])
    trm = mercado.niveles(corte)[C.FACTORES.index("TRM")]
    indice = mercado.niveles(corte)[C.FACTORES.index("COLCAP")]
    filas, canonicas = [], []
    for p in POSICIONES:
        if not _vigente(p, k, corte):
            continue
        errores = p.get("errores", {}).get(k, [])
        venc = pd.Timestamp(p["venc"]) if "venc" in p else None
        valorable = (venc is None or venc > corte) and normalizar_tipo(p["tipo"])[0] is not None
        fila = {"id": p["id"], "especie": p["especie"], "tipo": p["tipo"], "emisor": p["emisor"],
                "cal": p.get("cal", ""), "port": p["port"], "clas": p["clas"],
                "moneda": "USD" if "usd" in p else "COP", "nom": None, "tasa": p.get("tasa"),
                "margen": p.get("margen"), "modalidad": "EA" if ("tasa" in p or "margen" in p) else "",
                "per": p.get("per", ""), "emi": p.get("emi"), "venc": p.get("venc"),
                "vm": None, "beta": p.get("beta"), "restr": p.get("restr", "No"), "errores": errores}
        if "saldo" in p:
            fila["nom"] = fila["vm"] = p["saldo"][k] * 1e6
        elif "unidades" in p:
            fila["nom"] = p["unidades"]
            fila["vm"] = p["unidades"] * indice * (1 + rng.normal(0, 0.001))
        elif "usd" in p:
            fila["nom"] = p["usd"]
        else:
            fila["nom"] = p["nom"] * 1e6
        if valorable and "saldo" not in p and "unidades" not in p:
            canonicas.append((len(filas), _canonica(p, corte)))
        elif not valorable:
            fila["vm"] = fila["nom"] * 1.01
        filas.append(fila)

    # Valor de mercado "del proveedor": teórico con una diferencia pequeña de curva
    df = pd.DataFrame([c for _, c in canonicas])
    port = Portafolio(df, mercado, corte)
    valores = port.valorar(port.niveles0)[0]
    for (i, _), v in zip(canonicas, valores):
        filas[i]["vm"] = v * (1 + rng.uniform(-0.003, 0.003))
        if "valor_distinto" in filas[i]["errores"]:
            filas[i]["vm"] = v * 1.04
    return corte, filas, trm


def tabla_para_liquidez(corte, filas, mercado):
    """Arma la tabla valorada que usa el módulo de liquidez (para calibrar flujos)."""
    registros = []
    for f in filas:
        tipo, _ = normalizar_tipo(f["tipo"])
        if tipo is None or (f["venc"] and pd.Timestamp(f["venc"]) <= corte):
            continue
        cal = "NACION" if tipo == "TES_TF" else normalizar_calificacion(f["cal"])
        venc = pd.Timestamp(f["venc"]) if f["venc"] else pd.NaT
        registros.append({"id": f["id"], "especie": f["especie"], "tipo": tipo, "emisor": f["emisor"],
                          "valor": f["vm"], "restringido": f["restr"] == "Sí",
                          "plazo_anios": (venc - corte).days / 365 if pd.notna(venc) else np.nan,
                          "calificacion_lp": calificacion_lp(cal) if cal else None})
    # evitar doble conteo del duplicado inyectado (el motor lo elimina)
    return pd.DataFrame(registros).drop_duplicates("id")


def generar_flujos(corte, escala, rng_semilla):
    rng = np.random.default_rng(rng_semilla)
    filas = []
    dias_habiles = pd.bdate_range(corte + pd.Timedelta(days=1), corte + pd.Timedelta(days=120))
    for d in dias_habiles:
        dias = (d - corte).days
        if dias <= 90:
            peso = 1.6 if 16 <= dias <= 30 else 1.0 if dias <= 15 else 0.25
            filas.append((d, "Vencimiento de CDAT", "Salida", escala * peso * rng.uniform(0.6, 1.4) * 1e6))
        if d.weekday() == 3:
            filas.append((d, "Desembolsos de crédito comprometidos", "Salida", rng.uniform(4000, 5200) * 1e6))
            filas.append((d, "Recaudo de cartera", "Entrada", rng.uniform(6000, 7200) * 1e6))
    filas.append((corte + pd.Timedelta(days=12), "Pago de obligaciones financieras", "Salida", 2500e6))
    filas.append((corte + pd.Timedelta(days=40), "Pago de obligaciones financieras", "Salida", 3000e6))
    return pd.DataFrame(filas, columns=["fecha", "concepto", "sentido", "monto"]).sort_values("fecha")


def calibrar_flujos(k, corte, tabla, depositos, politica):
    """Busca la escala de vencimientos de CDAT que produce el IRL a 30 días objetivo."""
    bajo, alto = 10.0, 10000.0
    for _ in range(60):
        medio = (bajo + alto) / 2
        fl = generar_flujos(corte, medio, 100 + k)
        fl["dias"] = (fl["fecha"] - corte).dt.days
        fl["sentido"] = fl["sentido"].str.upper()
        irl = calcular_liquidez(tabla, fl, depositos, politica)["irl_30"]
        if irl > IRL30_OBJETIVO[k]:
            bajo = medio
        else:
            alto = medio
    return generar_flujos(corte, (bajo + alto) / 2, 100 + k)


COLS_INV = ["No.", "Especie / Nemotécnico", "Tipo de instrumento", "Emisor", "Calificación", "Portafolio",
            "Clasificación contable", "Moneda", "Valor nominal", "Tasa facial", "Margen IBR", "Modalidad",
            "Periodicidad", "Fecha emisión", "Fecha vencimiento", "Valor de mercado", "Beta", "Restringido"]
ANCHOS_INV = [6, 36, 20, 30, 11, 16, 22, 8, 17, 10, 10, 10, 11, 12, 13, 17, 6, 11]
FMT_INV = [None, None, None, None, None, None, None, None, FMT_MONTO, FMT_PCT, FMT_PCT, None, None,
           FMT_FECHA, FMT_FECHA, FMT_MONTO, "0.00", None]


def escribir_posiciones(ruta, k, corte, filas, flujos, depositos):
    wb = _libro()
    ws = wb.active
    ws.title = "Datos_generales"
    _titulos(ws, ENTIDAD.upper(), "Datos generales del corte")
    datos = [("Fecha de corte", corte.to_pydatetime(), FMT_FECHA),
             ("Patrimonio técnico (COP)", PATRIMONIO[k] * 1e6, FMT_MONTO),
             ("Elaborado por", "Tesorería", None), ("Revisado por", "Área de Riesgos", None)]
    for i, (a, b, fmt) in enumerate(datos):
        ws.cell(4 + i, 1, a).font = F_NEGRITA
        c = ws.cell(4 + i, 2, b)
        c.font = F_NORMAL
        if fmt:
            c.number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22

    ws = wb.create_sheet("Inversiones")
    _titulos(ws, ENTIDAD.upper(), f"Portafolio de inversiones - corte {corte:%d/%m/%Y}",
             "Cifras en pesos colombianos. Fuente: Tesorería (consolidado manual).")
    _encabezados(ws, 5, COLS_INV, ANCHOS_INV)
    r = 6
    for f in filas:
        e = f["errores"]
        cal = "" if "sin_calificacion" in e else ("AAA" if "calificacion_distinta" in e else f["cal"])
        tasa = f["tasa"] * 100 if ("tasa_pct" in e and f["tasa"]) else f["tasa"]
        venc = pd.Timestamp(f["venc"]).to_pydatetime() if f["venc"] else None
        if "venc_texto" in e:
            venc = pd.Timestamp(f["venc"]).strftime("%d/%m/%Y")
        nominal = f"{f['nom']:,.0f}".replace(",", ".") if "nominal_texto" in e else f["nom"]
        valores = [f["id"], f["especie"], f["tipo"], f["emisor"], cal, f["port"], f["clas"], f["moneda"],
                   nominal, tasa, f["margen"], f["modalidad"], f["per"],
                   pd.Timestamp(f["emi"]).to_pydatetime() if f["emi"] else None, venc,
                   round(f["vm"]), f["beta"], f["restr"]]
        _fila(ws, r, valores, FMT_INV)
        if "venc_texto" in e:
            ws.cell(r, 15).number_format = "@"
        if "nominal_texto" in e:
            ws.cell(r, 9).number_format = "@"
        r += 1
        if "duplicado" in e:
            _fila(ws, r, valores, FMT_INV)
            r += 1
    ultima = r - 1
    _fila(ws, r + 1, ["", "TOTAL", "", "", "", "", "", "", f"=SUM(I6:I{ultima})", "", "", "", "", "", "",
                      f"=SUM(P6:P{ultima})", "", ""], FMT_INV, negrita=True)
    ws.freeze_panes = "C6"

    ws = wb.create_sheet("Flujos_liquidez")
    _titulos(ws, ENTIDAD.upper(), "Flujos contractuales proyectados (próximos 120 días)")
    _encabezados(ws, 4, ["Fecha", "Concepto", "Sentido", "Monto (COP)"], [12, 38, 10, 18])
    extra = []
    if k == SEP:
        extra = [(corte - pd.Timedelta(days=2), "Vencimiento de CDAT (ya pagado)", "Salida", 850e6)]
    for i, (fe, con, sen, mon) in enumerate(list(flujos.itertuples(index=False)) + extra):
        _fila(ws, 5 + i, [pd.Timestamp(fe).to_pydatetime(), con, sen, round(mon)],
              [FMT_FECHA, None, None, FMT_MONTO])

    ws = wb.create_sheet("Depositos_vista")
    _titulos(ws, ENTIDAD.upper(), "Saldos de depósitos a la vista al corte")
    _encabezados(ws, 4, ["Producto", "Saldo (COP)"], [30, 20])
    for i, (prod, saldo) in enumerate(depositos.itertuples(index=False)):
        _fila(ws, 5 + i, [prod, round(saldo)], [None, FMT_MONTO])

    # El resumen que hoy se arma a mano: SUMAR.SI por el nombre tal como está escrito
    ws = wb.create_sheet("Resumen_manual")
    _titulos(ws, "Resumen de concentración por emisor (elaborado manualmente)",
             "Tabla dinámica / SUMAR.SI sobre la columna Emisor tal como está escrita.",
             "Tiempo estimado de elaboración mensual del informe manual: 4 horas.")
    _encabezados(ws, 5, ["Emisor", "Valor de mercado", "% del portafolio", "Límite", "Estado"],
                 [36, 18, 14, 10, 12])
    vistos, fila_r = [], 6
    for f in filas:
        clave = f["emisor"].strip().casefold()
        if clave in vistos:
            continue
        vistos.append(clave)
        es_nacion = "nación" in clave
        ws.cell(fila_r, 1, f["emisor"]).font = F_NORMAL
        ws.cell(fila_r, 2, f"=SUMIF(Inversiones!$D$6:$D${ultima},A{fila_r},Inversiones!$P$6:$P${ultima})")
        ws.cell(fila_r, 3, f"=IFERROR(B{fila_r}/$B${{T}},0)")
        ws.cell(fila_r, 4, "N.A." if es_nacion else 0.15)
        ws.cell(fila_r, 5, f'=IF(D{fila_r}="N.A.","N.A.",IF(C{fila_r}>D{fila_r},"EXCEDE","Cumple"))')
        for j, fmt in zip(range(1, 6), [None, FMT_MONTO, "0.0%", "0%", None]):
            ws.cell(fila_r, j).font = F_NORMAL
            if fmt:
                ws.cell(fila_r, j).number_format = fmt
        fila_r += 1
    total_r = fila_r + 1
    for rr in range(6, fila_r):
        ws.cell(rr, 3).value = ws.cell(rr, 3).value.replace("{T}", str(total_r))
    ws.cell(total_r, 1, "TOTAL").font = F_NEGRITA
    c = ws.cell(total_r, 2, f"=SUM(B6:B{fila_r - 1})")
    c.font, c.number_format = F_NEGRITA, FMT_MONTO
    wb.save(ruta)


def escribir_plantilla(ruta: Path):
    wb = _libro()
    ws = wb.active
    ws.title = "Datos_generales"
    _titulos(ws, "NOMBRE DE LA ENTIDAD", "Plantilla de carga mensual - diligencie las celdas amarillas")
    for i, (a, b, fmt) in enumerate([("Fecha de corte", "30/09/2026", None),
                                      ("Patrimonio técnico (COP)", 98_450_000_000, FMT_MONTO),
                                      ("Elaborado por", "", None), ("Revisado por", "", None)]):
        ws.cell(4 + i, 1, a).font = F_NEGRITA
        c = ws.cell(4 + i, 2, b)
        c.fill, c.font = RELLENO_EJEMPLO, F_NORMAL
        if fmt:
            c.number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22

    ws = wb.create_sheet("Inversiones")
    _titulos(ws, "NOMBRE DE LA ENTIDAD",
             "Una fila por título desde la fila 6. La fila amarilla es un EJEMPLO: reemplácela.",
             "Tipos: TES | CDT | CDT IBR | Bono ordinario | ETF | Depósito en USD | Cuenta de ahorros | FIC. "
             "Tasas en % (12,5%). Periodicidad: AV, SV, TV, MV o PV. Restringido: Sí/No. "
             "Nombre del archivo: posiciones_AAAA-MM-DD.xlsx")
    _encabezados(ws, 5, COLS_INV, ANCHOS_INV)
    ejemplo = [1, "TES TF 7,00% mar-2031", "TES", "Nación - MHCP", "Nación", "Tesorería",
               "Disponible para la venta", "COP", 1_000_000_000, 0.07, None, "EA", "AV",
               pd.Timestamp("2021-03-26").to_pydatetime(), pd.Timestamp("2031-03-26").to_pydatetime(),
               850_000_000, None, "No"]
    _fila(ws, 6, ejemplo, FMT_INV)
    for j in range(1, len(ejemplo) + 1):
        ws.cell(6, j).fill = RELLENO_EJEMPLO

    ws = wb.create_sheet("Flujos_liquidez")
    _titulos(ws, "NOMBRE DE LA ENTIDAD", "Flujos contractuales de los próximos 90 días (Sentido: Entrada/Salida)")
    _encabezados(ws, 4, ["Fecha", "Concepto", "Sentido", "Monto (COP)"], [12, 38, 10, 18])
    _fila(ws, 5, [pd.Timestamp("2026-10-07").to_pydatetime(), "Vencimiento de CDAT", "Salida", 1_500_000_000],
          [FMT_FECHA, None, None, FMT_MONTO])
    for j in range(1, 5):
        ws.cell(5, j).fill = RELLENO_EJEMPLO

    ws = wb.create_sheet("Depositos_vista")
    _titulos(ws, "NOMBRE DE LA ENTIDAD", "Saldos de depósitos a la vista al corte")
    _encabezados(ws, 4, ["Producto", "Saldo (COP)"], [30, 20])
    _fila(ws, 5, ["Ahorro a la vista", 150_000_000_000], [None, FMT_MONTO])
    for j in range(1, 3):
        ws.cell(5, j).fill = RELLENO_EJEMPLO
    wb.save(ruta)


def main():
    rng = np.random.default_rng(7)
    (RAIZ / "mercado").mkdir(exist_ok=True)
    (RAIZ / "entrada" / "plantilla").mkdir(parents=True, exist_ok=True)

    print("1/4 Mercado sintético ...")
    mercado_sintetico.generar().to_csv(C.ARCHIVO_MERCADO, encoding="utf-8-sig")
    mercado = Mercado(C.ARCHIVO_MERCADO)

    print("2/4 Política de riesgos ...")
    escribir_politica(C.ARCHIVO_POLITICA)
    politica = leer_politica(C.ARCHIVO_POLITICA)

    print("3/4 Posiciones mensuales ...")
    for k in range(len(CORTES)):
        corte, filas, _ = construir_corte(k, mercado, rng)
        depositos = pd.DataFrame({"producto": ["Ahorro a la vista", "Ahorro programado"],
                                  "saldo": [(142_000 + 1_500 * k) * 1e6, (21_000 + 300 * k) * 1e6]})
        tabla = tabla_para_liquidez(corte, filas, mercado)
        flujos = calibrar_flujos(k, corte, tabla, depositos, politica)
        ruta = C.DIR_ENTRADA / f"posiciones_{CORTES[k]}.xlsx"
        escribir_posiciones(ruta, k, corte, filas, flujos, depositos)
        print(f"    {ruta.name}")

    print("4/4 Plantilla vacía ...")
    escribir_plantilla(C.DIR_ENTRADA / "plantilla" / "plantilla_posiciones.xlsx")
    print("Listo. Siguiente paso: python generar_reporte.py --todos")


if __name__ == "__main__":
    main()
