"""Anexo en Excel: el detalle de todo el informe, para auditoría y revisoría.

Las cifras que produce la simulación (VaR, ES, escenarios) se escriben como
valores; las que son aritmética simple se escriben como FÓRMULAS vivas
(totales, pesos, excepciones del backtesting, cadena del IRL), para que
cualquier revisor pueda rastrear cómo se obtiene cada número.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from . import config as C
from .liquidez import DESCRIPCION_CATEGORIA

FUENTE = "Arial"
F_TIT = Font(name=FUENTE, size=12, bold=True, color="1F3B5C")
F_SUB = Font(name=FUENTE, size=9, italic=True, color="52514E")
F_ENC = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
F_N = Font(name=FUENTE, size=9)
F_B = Font(name=FUENTE, size=9, bold=True)
R_ENC = PatternFill("solid", fgColor="1F3B5C")
R_TOTAL = PatternFill("solid", fgColor="EEF1F5")
BORDE = Border(bottom=Side(style="thin", color="D9DCE1"))
ESTADO_RELLENO = {"ROJO": "D03B3B", "ÁMBAR": "FAB219", "VERDE": "0CA30C", "SIN DATO": "C3C2B7"}
ESTADO_FUENTE = {"ROJO": "FFFFFF", "ÁMBAR": "0B0B0B", "VERDE": "0B0B0B", "SIN DATO": "0B0B0B"}

M_COP = '#,##0;(#,##0);"-"'
M_MILL = '#,##0.0;(#,##0.0);"-"'
M_PCT = '0.0%;(0.0%);"-"'
M_PCT2 = '0.00%;(0.00%);"-"'
M_DEC = '0.00;(0.00);"-"'
M_FECHA = "dd/mm/yyyy"
FORMATO_UNIDAD = {"%": M_PCT, "COP M": M_MILL, "años": M_DEC, "número": "0", "puntos": "#,##0"}


def _titulo(ws, titulo, sub=None):
    ws["A1"] = titulo
    ws["A1"].font = F_TIT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def _encabezado(ws, fila, columnas, anchos=None):
    for j, nombre in enumerate(columnas, start=1):
        c = ws.cell(fila, j, nombre)
        c.font, c.fill = F_ENC, R_ENC
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if anchos:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 30


def _escribir(ws, fila, valores, formatos=None, negrita=False, relleno=None):
    for j, v in enumerate(valores, start=1):
        if isinstance(v, (np.floating, np.integer)):
            v = v.item()
        if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
            v = None
        if isinstance(v, pd.Timestamp):
            v = None if pd.isna(v) else v.to_pydatetime()
        c = ws.cell(fila, j, v)
        c.font = F_B if negrita else F_N
        c.border = BORDE
        if relleno:
            c.fill = relleno
        if formatos and j - 1 < len(formatos) and formatos[j - 1]:
            c.number_format = formatos[j - 1]


def _estado(celda, estado):
    celda.fill = PatternFill("solid", fgColor=ESTADO_RELLENO[estado])
    celda.font = Font(name=FUENTE, size=9, bold=True, color=ESTADO_FUENTE[estado])
    celda.alignment = Alignment(horizontal="center")


# ------------------------------------------------------------------- hojas
def _resumen(wb, r):
    ws = wb.active
    ws.title = "Resumen"
    ind = r["indicadores"]
    _titulo(ws, f"{r['politica']['nombre']} - Anexo del informe de riesgos",
            f"Corte {r['corte']:%d/%m/%Y}. Cifras en COP millones salvo indicación.")
    kpis = [("Valor del portafolio (COP M)", ind["VALOR_PORTAFOLIO"], M_MILL),
            ("Patrimonio técnico (COP M)", ind["PATRIMONIO_TECNICO"], M_MILL),
            ("VaR histórico 99% 1 día (COP M)", ind["VAR_HS_1D"], M_MILL),
            ("ES histórico 97,5% 1 día (COP M)", ind["ES_HS_1D"], M_MILL),
            ("VaR paramétrico 99% 1 día (COP M)", ind["VAR_PARAM_1D"], M_MILL),
            ("IRL razón 30 días", ind["IRL_30D"], M_PCT)]
    for i, (k, v, fmt) in enumerate(kpis):
        ws.cell(4 + i, 1, k).font = F_B
        _escribir(ws, 4 + i, [k, v], [None, fmt])
        ws.cell(4 + i, 1).font = F_B
    fila0 = 4 + len(kpis) + 2
    _encabezado(ws, fila0, ["Código", "Indicador", "Categoría", "Tipo", "Valor", "Límite", "Alerta",
                            "Estado", "Detalle"], [26, 52, 16, 10, 12, 12, 12, 12, 34])
    for i, (_, l) in enumerate(r["limites"].iterrows(), start=1):
        fmt = FORMATO_UNIDAD.get(l["unidad"], M_DEC)
        _escribir(ws, fila0 + i, [l["codigo"], l["indicador"], l["categoria"], l["tipo"], l["valor"],
                                  l["limite"], l["alerta"], l["estado"].capitalize(), l["detalle"]],
                  [None, None, None, None, fmt, fmt, fmt, None, None])
        _estado(ws.cell(fila0 + i, 8), l["estado"])
    ws.freeze_panes = ws.cell(fila0 + 1, 1)
    fila_m = fila0 + len(r["limites"]) + 3
    ws.cell(fila_m, 1, "Mensajes clave").font = F_TIT
    for i, (tono, texto) in enumerate(r["mensajes"], start=1):
        ws.cell(fila_m + i, 1, {"rojo": "Excedido", "ambar": "Alerta", "info": "Info"}[tono]).font = F_B
        c = ws.cell(fila_m + i, 2, texto)
        c.font, c.alignment = F_N, Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila_m + i, start_column=2, end_row=fila_m + i, end_column=9)
        ws.row_dimensions[fila_m + i].height = 15 * max(1, len(texto) // 150 + 1)
    fila_a = fila_m + len(r["mensajes"]) + 3
    ws.cell(fila_a, 1, "Acciones sugeridas").font = F_TIT
    _encabezado(ws, fila_a + 1, ["Prioridad", "Tema", "Acción sugerida", "", "", "", "", "", "Monto (COP M)"])
    for i, a in enumerate(r["acciones"], start=2):
        f = fila_a + i
        _escribir(ws, f, [a["prioridad"].capitalize(), a["tema"], a["accion"], None, None, None, None, None,
                          None if a["monto"] is None else a["monto"] / 1e6],
                  [None] * 8 + [M_MILL])
        if a["prioridad"] in ESTADO_RELLENO:
            _estado(ws.cell(f, 1), a["prioridad"])
        ws.cell(f, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=8)
        ws.row_dimensions[f].height = 15 * max(1, len(a["accion"]) // 120 + 1)


def _posiciones(wb, r):
    ws = wb.create_sheet("Posiciones")
    _titulo(ws, "Detalle de posiciones valoradas",
            "Valor teórico = valoración del motor. Diferencia = valor reportado por la entidad / teórico - 1.")
    cols = ["ID", "Especie", "Tipo", "Emisor unificado", "Nombre en el archivo", "Sector", "Calificación",
            "Portafolio", "Clasificación contable", "Moneda", "Valor nominal", "Tasa facial", "Margen IBR",
            "Vencimiento", "Plazo (años)", "Valor teórico (COP)", "Valor reportado (COP)", "Diferencia",
            "Peso", "Duración mod.", "DV01 (COP por pb)", "TIR", "Contribución ES (COP)",
            "Categoría liquidez", "Haircut", "Computable ALAC (COP)", "Fila en el Excel origen"]
    anchos = [7, 36, 15, 30, 30, 14, 11, 16, 20, 8, 16, 9, 9, 11, 9, 17, 17, 10, 8, 10, 14, 8, 16, 26, 8, 17, 11]
    _encabezado(ws, 4, cols, anchos)
    t = r["tabla"]
    n = len(t)
    fin = 4 + n
    fmts = [None, None, None, None, None, None, None, None, None, None, M_COP, M_PCT2, M_PCT2, M_FECHA, M_DEC,
            M_COP, M_COP, M_PCT, M_PCT, M_DEC, M_COP, M_PCT2, M_COP, None, M_PCT, M_COP, "0"]
    for i, (_, f) in enumerate(t.iterrows()):
        fila = 5 + i
        _escribir(ws, fila, [f["id"], f["especie"], C.NOMBRE_TIPO[f["tipo"]], f["emisor"], f["emisor_original"],
                             f["sector"], f["calificacion"], f["portafolio"], f["clasificacion"], f["moneda"],
                             f["valor_nominal"], f["tasa_facial"], f["margen"], f["fecha_vencimiento"],
                             f["plazo_anios"], f["valor"],
                             f["valor_mercado"] if pd.notna(f["valor_mercado"]) else None,
                             f'=IF(Q{fila}="","",Q{fila}/P{fila}-1)', f"=P{fila}/$P${fin + 1}",
                             f["duracion_mod"], f["dv01"], f["tir"], f["contrib_es"],
                             DESCRIPCION_CATEGORIA.get(f["categoria"], f["categoria"]), f["haircut"],
                             f"=P{fila}*(1-Y{fila})", f["fila_excel"]], fmts)
    _escribir(ws, fin + 1, ["", "TOTAL"] + [""] * 13 + [f"=SUM(P5:P{fin})", f"=SUM(Q5:Q{fin})", "",
                                                        f"=SUM(S5:S{fin})", "", f"=SUM(U5:U{fin})", "",
                                                        f"=SUM(W5:W{fin})", "", "", f"=SUM(Z5:Z{fin})", ""],
              fmts, negrita=True, relleno=R_TOTAL)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{fin}"
    return fin + 1   # fila de totales (la usa la hoja de liquidez)


def _mercado(wb, r):
    ws = wb.create_sheet("Riesgo_mercado")
    v = r["var"]
    _titulo(ws, "VaR y Expected Shortfall",
            "Los porcentajes son fórmulas sobre el valor del portafolio y el patrimonio técnico (celdas B4 y B5).")
    ws["A4"], ws["B4"] = "Valor del portafolio (COP)", v["valor_total"]
    ws["A5"], ws["B5"] = "Patrimonio técnico (COP)", r["patrimonio_tecnico"]
    for c in ("A4", "A5"):
        ws[c].font = F_B
    for c in ("B4", "B5"):
        ws[c].font, ws[c].number_format = F_N, M_COP
    _encabezado(ws, 7, ["Medida (1 día)", "COP", "% portafolio", "% patrimonio técnico", "A 10 días (COP)"],
                [44, 18, 14, 18, 18])
    medidas = [("VaR 99% histórico", v["var_hs"]), ("ES 97,5% histórico", v["es_hs"]),
               ("VaR 99% paramétrico EWMA", v["var_param"]), ("ES 97,5% paramétrico normal", v["es_param"])]
    for i, (nombre, valor) in enumerate(medidas):
        f = 8 + i
        _escribir(ws, f, [nombre, valor, f"=B{f}/$B$4", f"=B{f}/$B$5",
                          f"=B{f}*SQRT({C.DIAS_HORIZONTE_REGULATORIO})"], [None, M_COP, M_PCT2, M_PCT2, M_COP])
    _encabezado(ws, 14, ["Factor de riesgo", "VaR histórico (COP)", "VaR paramétrico (COP)"])
    for i, g in enumerate(C.GRUPOS_FACTORES):
        _escribir(ws, 15 + i, [g, v["grupos_hs"][g], v["grupos_param"][g]], [None, M_COP, M_COP])
    f = 15 + len(C.GRUPOS_FACTORES)
    _escribir(ws, f, ["Suma de VaR individuales", f"=SUM(B15:B{f - 1})", f"=SUM(C15:C{f - 1})"],
              [None, M_COP, M_COP], negrita=True)
    _escribir(ws, f + 1, ["VaR total (diversificado)", v["var_hs"], v["var_param"]], [None, M_COP, M_COP],
              negrita=True)
    _escribir(ws, f + 2, ["Beneficio de diversificación", f"=B{f}-B{f + 1}", f"=C{f}-C{f + 1}"],
              [None, M_COP, M_COP], negrita=True, relleno=R_TOTAL)
    _encabezado(ws, f + 5, ["Peores escenarios históricos", "P&G (COP)"])
    for i, (fecha, pnl) in enumerate(v["peores_dias"]):
        _escribir(ws, f + 6 + i, [fecha, pnl], [M_FECHA, M_COP])


def _sensibilidades(wb, r):
    ws = wb.create_sheet("Sensibilidades")
    _titulo(ws, "Sensibilidades por posición",
            "DV01 por nodo = pérdida en COP si ese nodo de la curva TES sube 1 pb. Deltas = cambio en COP "
            "ante +1% en la TRM o en el índice accionario.")
    t = r["tabla"]
    nodos = list(C.NODOS_CURVA)
    cols = ["ID", "Especie"] + [f"DV01 {C.ETIQUETA_NODO[n]}" for n in nodos] + \
           ["DV01 curva", "DV01 spread", "Delta TRM +1%", "Delta RV +1%"]
    _encabezado(ws, 4, cols, [7, 36] + [11] * len(nodos) + [13, 13, 14, 14])
    k = len(nodos)
    for i, (_, f) in enumerate(t.iterrows()):
        fila = 5 + i
        primera, ultima = get_column_letter(3), get_column_letter(2 + k)
        _escribir(ws, fila, [f["id"], f["especie"]] + [f[f"krd_{n}"] for n in nodos] +
                  [f"=SUM({primera}{fila}:{ultima}{fila})", f["dv01_spread"], f["delta_trm_1pct"],
                   f["delta_rv_1pct"]], [None, None] + [M_COP] * (k + 4))
    fin = 4 + len(t)
    total = ["", "TOTAL"] + [f"=SUM({get_column_letter(j)}5:{get_column_letter(j)}{fin})"
                             for j in range(3, 3 + k + 4)]
    _escribir(ws, fin + 1, total, [None, None] + [M_COP] * (k + 4), negrita=True, relleno=R_TOTAL)
    ws.freeze_panes = "C5"


def _backtesting(wb, r):
    ws = wb.create_sheet("Backtesting")
    bt = r["backtesting"]
    _titulo(ws, "Backtesting hipotético",
            "Excepción = 1 si la pérdida del día supera el VaR calculado con información del día anterior "
            "(columnas E y F son fórmulas).")
    _encabezado(ws, 4, ["Prueba", "VaR histórico", "VaR paramétrico"], [30, 16, 16])
    filas = [("Excepciones", "=SUM(E11:E{})", "=SUM(F11:F{})"),
             ("Esperadas", bt["hs"]["esperadas"], bt["param"]["esperadas"]),
             ("Zona de Basilea", bt["hs"]["zona"], bt["param"]["zona"]),
             ("Kupiec (p-valor)", bt["hs"]["kupiec_p"], bt["param"]["kupiec_p"]),
             ("Christoffersen (p-valor)", bt["hs"]["christoffersen_p"], bt["param"]["christoffersen_p"])]
    fin = 10 + bt["n"]
    for i, (a, b, c) in enumerate(filas):
        b = b.format(fin) if isinstance(b, str) and "{}" in b else b
        c = c.format(fin) if isinstance(c, str) and "{}" in c else c
        _escribir(ws, 5 + i, [a, b, c], [None, "0.000" if i >= 3 else "0.0", "0.000" if i >= 3 else "0.0"])
    _encabezado(ws, 10, ["Fecha", "P&G hipotético (COP)", "VaR histórico (COP)", "VaR paramétrico (COP)",
                         "Excepción histórico", "Excepción paramétrico"], [30, 16, 16, 18, 14, 14])
    for i, fecha in enumerate(bt["fechas"]):
        f = 11 + i
        _escribir(ws, f, [fecha, bt["pnl"][i], bt["var_hs"][i], bt["var_param"][i],
                          f"=IF(B{f}<-C{f},1,0)", f"=IF(B{f}<-D{f},1,0)"],
                  [M_FECHA, M_COP, M_COP, M_COP, "0", "0"])
    ws.freeze_panes = "A11"


def _estres(wb, r):
    ws = wb.create_sheet("Estres")
    _titulo(ws, "Pruebas de estrés", "% del patrimonio técnico = fórmula sobre la celda B3.")
    ws["A3"], ws["B3"] = "Patrimonio técnico (COP)", r["patrimonio_tecnico"]
    ws["A3"].font, ws["B3"].font, ws["B3"].number_format = F_B, F_N, M_COP
    _encabezado(ws, 5, ["Código", "Escenario", "Tipo", "P&G (COP)", "% patrimonio técnico", "Choque aplicado",
                        "Ventana histórica", "Posición más afectada", "P&G de esa posición (COP)"],
                [8, 36, 11, 16, 12, 70, 24, 36, 18])
    for i, (_, e) in enumerate(r["estres"].iterrows()):
        f = 6 + i
        _escribir(ws, f, [e["codigo"], e["escenario"], e["tipo"], e["pnl"], f"=D{f}/$B$3", e["choque"],
                          e["ventana"], e["posicion_mas_afectada"], e["pnl_posicion_mas_afectada"]],
                  [None, None, None, M_COP, M_PCT2, None, None, None, M_COP])


def _concentracion(wb, r):
    ws = wb.create_sheet("Concentracion")
    c = r["concentracion"]
    _titulo(ws, "Concentración por emisor unificado",
            "'Nombres en el archivo' muestra todas las formas en que el emisor aparece escrito.")
    _encabezado(ws, 4, ["Emisor unificado", "Nombres en el archivo", "Sector", "Calificación", "Exento",
                        "Títulos", "Valor (COP)", "% portafolio", "% patrimonio técnico"],
                [34, 60, 14, 12, 8, 8, 17, 12, 14])
    e = c["emisores"]
    fin = 4 + len(e)
    for i, (_, f) in enumerate(e.iterrows()):
        fila = 5 + i
        _escribir(ws, fila, [f["emisor"], " | ".join(f["nombres"]), f["sector"], f["calificacion"],
                             "Sí" if f["exento"] else "No", f["titulos"], f["valor"],
                             f"=G{fila}/$G${fin + 1}", f["pct_pt"]], [None] * 5 + ["0", M_COP, M_PCT, M_PCT])
    _escribir(ws, fin + 1, ["TOTAL", "", "", "", "", f"=SUM(F5:F{fin})", f"=SUM(G5:G{fin})",
                            f"=SUM(H5:H{fin})", ""], [None] * 5 + ["0", M_COP, M_PCT, M_PCT],
              negrita=True, relleno=R_TOTAL)
    fila = fin + 4
    for titulo, df in (("Por sector", c["por_sector"]), ("Por calificación", c["por_calificacion"]),
                       ("Por plazo", c["por_plazo"]), ("Por tipo de instrumento", c["por_tipo"])):
        _encabezado(ws, fila, [titulo, "Valor (COP)", "% portafolio"])
        for i, (_, g) in enumerate(df.iterrows(), start=1):
            _escribir(ws, fila + i, [g["grupo"], g["valor"], g["pct"]], [None, M_COP, M_PCT])
        fila += len(df) + 3
    _escribir(ws, fila, ["HHI emisores privados", c["hhi_privados"]], [None, "#,##0"], negrita=True)
    _escribir(ws, fila + 1, ["Número equivalente de emisores", c["n_equivalente"]], [None, "0.0"], negrita=True)


def _liquidez(wb, r, fila_total_pos):
    ws = wb.create_sheet("Liquidez")
    liq = r["liquidez"]
    tope = r["politica"]["liquidez"].get("tope de entradas", 0.75)
    _titulo(ws, "IRL simplificado por bandas",
            "Cadena completa en fórmulas: entradas computables = MIN(entradas; tope x salidas); "
            "RLN = salidas - entradas computables; IRL = ALAC / RLN acumulado.")
    ws["A3"], ws["B3"] = "ALAC (COP) - suma de la hoja Posiciones", f"=Posiciones!Z{fila_total_pos}"
    ws["A4"], ws["B4"] = "Tope de entradas", tope
    for cel, fmt in (("B3", M_COP), ("B4", M_PCT)):
        ws[cel].number_format, ws[cel].font = fmt, F_N
    ws["A3"].font = ws["A4"].font = F_B
    _encabezado(ws, 6, ["Banda", "Salidas contractuales (COP)", "Retiros estimados (COP)", "Entradas (COP)",
                        "Entradas computables (COP)", "RLN (COP)", "RLN acumulado (COP)", "IRL monto (COP)",
                        "IRL razón"], [22, 18, 18, 16, 18, 16, 18, 18, 11])
    for i, (_, b) in enumerate(liq["bandas"].iterrows()):
        f = 7 + i
        _escribir(ws, f, [b["banda"], b["salidas_contractuales"], b["retiros_estimados"], b["entradas"],
                          f"=MIN(D{f},$B$4*(B{f}+C{f}))", f"=B{f}+C{f}-E{f}",
                          f"=F{f}" if i == 0 else f"=G{f - 1}+F{f}", f"=$B$3-G{f}", f"=IFERROR($B$3/G{f},\"\")"],
                  [None] + [M_COP] * 7 + [M_PCT])
    _encabezado(ws, 13, ["Flujos contractuales usados", "Concepto", "Sentido", "Monto (COP)", "Días al pago"])
    for i, (_, fl) in enumerate(r["flujos"].sort_values("fecha").iterrows()):
        _escribir(ws, 14 + i, [fl["fecha"], fl["concepto"], fl["sentido"], fl["monto"], fl["dias"]],
                  [M_FECHA, None, None, M_COP, "0"])


def _calidad(wb, r):
    ws = wb.create_sheet("Calidad_datos")
    _titulo(ws, "Bitácora de calidad de datos", "Hoja y fila exactas del archivo de la entidad.")
    _encabezado(ws, 4, ["Severidad", "Hoja", "Fila Excel", "Campo", "Valor original", "Detalle", "Acción"],
                [13, 16, 9, 18, 30, 70, 50])
    for i, (_, h) in enumerate(r["hallazgos"].tabla().iterrows()):
        _escribir(ws, 5 + i, list(h.values), [None, None, "0", None, None, None, None])
        estado = {"ERROR": "ROJO", "ADVERTENCIA": "ÁMBAR", "INFO": "SIN DATO"}[h["severidad"]]
        _estado(ws.cell(5 + i, 1), estado)
    ws.freeze_panes = "A5"


def _parametros(wb, r):
    ws = wb.create_sheet("Parametros")
    tr = r["traza"]
    _titulo(ws, "Parámetros, supuestos y trazabilidad",
            "Supuestos metodológicos del motor y huella SHA-256 de cada archivo de entrada.")
    _encabezado(ws, 4, ["Parámetro", "Valor"], [44, 90])
    filas = list(tr["parametros"].items()) + [
        ("Versión del motor", tr["version"]), ("Fecha y hora de ejecución", tr["ejecutado"]),
        ("Fuente de la curva y spreads", "Archivo mercado/factores_mercado.csv (en producción: proveedor de precios)"),
        ("Supuesto de valoración", "Descuento con curva cero cupón TES + spread por calificación del emisor"),
        ("Supuesto de flotantes IBR", "Próximo cupón fijado; después, nominal + VP de (margen pactado - spread)"),
        ("Supuesto del IRL", "Metodología simplificada e ilustrativa; ajustar al estándar del supervisor")]
    for i, (k, v) in enumerate(filas):
        _escribir(ws, 5 + i, [k, v])
    fila = 7 + len(filas)
    _encabezado(ws, fila, ["Archivo de entrada", "Huella SHA-256"])
    for i, (nombre, huella) in enumerate(tr["archivos"], start=1):
        _escribir(ws, fila + i, [nombre, huella])


def generar_anexo(r: dict, ruta: Path) -> Path:
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    _resumen(wb, r)
    fila_total = _posiciones(wb, r)
    _mercado(wb, r)
    _sensibilidades(wb, r)
    _backtesting(wb, r)
    _estres(wb, r)
    _concentracion(wb, r)
    _liquidez(wb, r, fila_total)
    _calidad(wb, r)
    _parametros(wb, r)
    wb.save(ruta)
    return ruta
